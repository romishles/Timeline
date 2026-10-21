"""
Throwaway copies of the repository, and a small runner for the GitHub
workflow files, for the tests in this folder.

Every test works in its own temporary folder:

  origin.git  a bare repository standing in for GitHub (pushes go here)
  work/       your clone: edit files, commit and push like an upload
  runner-N/   a fresh clone per workflow run, like actions/checkout

The real repository is only ever read (to copy its tracked files once per
test run). Large media files are copied as small placeholders: the scripts
only check that they exist.

WorkflowRun runs the `run:` steps of a workflow file in order, with their
`if:`, `env:` and `${{ }}` expressions, GITHUB_OUTPUT, GITHUB_STEP_SUMMARY
and RUNNER_TEMP, the way GitHub Actions does. Steps with `uses:` (checkout,
setup-python) are replaced by the clone and the local Python. `pip` is a
no-op and `gh` always fails (as the Pages request does in the real repo).
"""

import json
import os
import re
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

REPO = Path(__file__).resolve().parent.parent

# Copied as small placeholder files: the build and the cleanup only check
# that these exist, and the real ones are about 200 MB.
PLACEHOLDER = re.compile(r"^(assets/(images|videos)/.*\.(jpg|mp4)|(video|image)-originals/[0-9]+\.\w+)$")
# What the scripts and workflows need; the site's own code is left out.
NEEDED = ("data/", "assets/", "xlsx_to_sqlite/", ".github/", "video-originals/", "image-originals/")

_template = None
_tmp_root = None


def run(cmd, cwd, env=None, check=True):
    result = subprocess.run(cmd, cwd=cwd, env=env, capture_output=True, text=True)
    if check and result.returncode != 0:
        raise RuntimeError(f"{' '.join(map(str, cmd))} failed in {cwd}:\n{result.stdout}\n{result.stderr}")
    return result


def git_env(home):
    env = {k: v for k, v in os.environ.items() if not k.startswith("GIT_")}
    env.update(GIT_CONFIG_NOSYSTEM="1", GIT_CONFIG_GLOBAL=str(home / "gitconfig"), HOME=str(home),
               PYTHONDONTWRITEBYTECODE="1")
    return env


def tmp_root():
    global _tmp_root
    if _tmp_root is None:
        _tmp_root = Path(tempfile.mkdtemp(prefix="timeline-tests-"))
    return _tmp_root


def cleanup():
    if _tmp_root is not None and not os.environ.get("KEEP_SANDBOX"):
        shutil.rmtree(_tmp_root, ignore_errors=True)


def template():
    """A bare repository with one commit holding the real repo's tracked files
    as they are in the working tree (so uncommitted changes are tested too)."""
    global _template
    if _template:
        return _template
    root = tmp_root()
    home = root / "home"
    home.mkdir()
    (home / "gitconfig").write_text("[init]\n\tdefaultBranch = main\n[user]\n\tname = Test\n\temail = test@example.com\n"
                                    "[advice]\n\tdetachedHead = false\n")
    env = git_env(home)
    src = root / "template-src"
    files = run(["git", "ls-files", "-z"], REPO).stdout.split("\0")
    for name in filter(None, files):
        if not name.startswith(NEEDED) or not (REPO / name).is_file():
            continue
        target = src / name
        target.parent.mkdir(parents=True, exist_ok=True)
        if PLACEHOLDER.match(name):
            target.write_text(f"placeholder for {name}\n")
        else:
            shutil.copyfile(REPO / name, target)
    run(["git", "init", "-q"], src, env)
    run(["git", "add", "-A"], src, env)
    run(["git", "commit", "-q", "-m", "Start"], src, env)
    _template = root / "template.git"
    run(["git", "clone", "-q", "--bare", str(src), str(_template)], root, env)
    shutil.rmtree(src)
    return _template


class Sandbox:
    """One throwaway origin plus a working clone."""

    count = 0

    def __init__(self):
        Sandbox.count += 1
        self.root = tmp_root() / f"sandbox-{Sandbox.count}"
        self.root.mkdir()
        self.home = tmp_root() / "home"
        self.env = git_env(self.home)
        self.origin = self.root / "origin.git"
        run(["git", "clone", "-q", "--bare", str(template()), str(self.origin)], self.root, self.env)
        self.work = self.root / "work"
        run(["git", "clone", "-q", str(self.origin), str(self.work)], self.root, self.env)
        self.runs = 0

    # --- git -----------------------------------------------------------------
    def git(self, *args, cwd=None, check=True):
        return run(["git", *args], cwd or self.work, self.env, check)

    def push(self, message):
        """Commits everything in work/ and pushes it, like an upload to main."""
        self.git("add", "-A")
        self.git("commit", "-q", "-m", message)
        self.git("push", "-q", "origin", "HEAD:main")

    def pull(self):
        self.git("pull", "-q", "--rebase", "origin", "main")

    def origin_head(self):
        return self.git("rev-parse", "main", cwd=self.origin).stdout.strip()

    def origin_subjects(self, since):
        """Subjects of the commits on origin's main after `since`, oldest first."""
        out = self.git("log", "--reverse", "--format=%s", f"{since}..main", cwd=self.origin).stdout
        return out.splitlines()

    def origin_show(self, path, ref="main"):
        return self.git("show", f"{ref}:{path}", cwd=self.origin).stdout

    def origin_bytes(self, path, ref="main"):
        return subprocess.run(["git", "show", f"{ref}:{path}"], cwd=self.origin, env=self.env,
                              capture_output=True, check=True).stdout

    def origin_has(self, path, ref="main"):
        return self.git("cat-file", "-e", f"{ref}:{path}", cwd=self.origin, check=False).returncode == 0

    # --- editing the working clone ---------------------------------------------
    def path(self, name):
        return self.work / name

    def edit_workbook(self, change):
        """Opens data/timeline.xlsx in work/, calls change(sheet), saves it."""
        from openpyxl import load_workbook
        path = self.path("data/timeline.xlsx")
        workbook = load_workbook(path)
        change(workbook["Sheet1"])
        workbook.save(path)

    def delete_items(self, ids, clear_links=True):
        """Deletes the rows of these ids, and (like a careful editor) clears link
        cells of other rows that point at them."""
        ids = set(ids)

        def change(sheet):
            header = [c.value for c in sheet[1]]
            link_cols = [i + 1 for i, h in enumerate(header) if isinstance(h, str) and h.startswith("link")]
            for row in range(sheet.max_row, 1, -1):
                if sheet.cell(row, 1).value in ids:
                    sheet.delete_rows(row)
            if clear_links:
                for row in range(2, sheet.max_row + 1):
                    for col in link_cols:
                        if sheet.cell(row, col).value in ids:
                            sheet.cell(row, col).value = None
        self.edit_workbook(change)

    def set_cell(self, item_id, column, value):
        def change(sheet):
            header = [c.value for c in sheet[1]]
            col = header.index(column) + 1
            for row in range(2, sheet.max_row + 1):
                if sheet.cell(row, 1).value == item_id:
                    sheet.cell(row, col).value = value
                    return
            raise KeyError(item_id)
        self.edit_workbook(change)

    def read_json(self, name):
        return json.loads(self.path(name).read_text(encoding="utf-8"))

    def write_json(self, name, data):
        self.path(name).write_text(json.dumps(data, indent=2), encoding="utf-8")

    # --- running the tools and the workflows ------------------------------------
    def build(self, *extra, output="new.db", baseline=True):
        """Runs the converter in work/ the way the workflow does."""
        cmd = [sys.executable, "xlsx_to_sqlite/xlsx_to_sqlite.py", "data/timeline.xlsx", output,
               "--videos", "data/videos.csv", "--manifest", "assets/videos/manifest.json",
               "--images", "assets/images/manifest.json"]
        if baseline:
            cmd += ["--baseline", "assets/timeline_data.db"]
        return run(cmd + list(extra), self.work, self.env, check=False)

    def workflow(self, file, event="push", inputs=None):
        self.runs += 1
        return WorkflowRun(self, file, event, inputs or {}, self.root / f"runner-{self.runs}").run()


# --- a small YAML reader, enough for the two workflow files ----------------------

def parse_yaml(text):
    lines = text.split("\n")

    def scalar(value):
        value = value.strip()
        if value[:1] in "'\"" and value[-1:] == value[:1]:
            return value[1:-1]
        if value.startswith("[") and value.endswith("]"):
            return [scalar(v) for v in value[1:-1].split(",") if v.strip()]
        return {"true": True, "false": False}.get(value, value)

    def significant(i):
        while i < len(lines) and (not lines[i].strip() or lines[i].lstrip().startswith("#")):
            i += 1
        return i

    def indent_of(line):
        return len(line) - len(line.lstrip(" "))

    def block_scalar(i, parent_indent):
        body = []
        while i < len(lines) and (not lines[i].strip() or indent_of(lines[i]) > parent_indent):
            body.append(lines[i])
            i += 1
        while body and not body[-1].strip():
            body.pop()
        cut = min(indent_of(l) for l in body if l.strip())
        return "\n".join(l[cut:] for l in body) + "\n", i

    def parse(i, indent):
        i = significant(i)
        if i >= len(lines):
            return None, i
        if lines[i].strip().startswith("- "):
            items = []
            while True:
                i = significant(i)
                if i >= len(lines) or indent_of(lines[i]) != indent or not lines[i].strip().startswith("- "):
                    return items, i
                # "- key: value" starts a mapping indented by two more
                lines[i] = " " * (indent + 2) + lines[i].strip()[2:]
                item, i = parse(i, indent + 2)
                items.append(item)
        mapping = {}
        while True:
            i = significant(i)
            if i >= len(lines) or indent_of(lines[i]) != indent:
                return mapping, i
            key, _, value = lines[i].strip().partition(":")
            value = re.sub(r"\s+#.*$", "", value) if not value.strip().startswith(("'", '"')) else value
            if value.strip() == "|":
                mapping[key], i = block_scalar(i + 1, indent)
            elif value.strip():
                mapping[key] = scalar(value)
                i += 1
            else:
                nxt = significant(i + 1)
                if nxt < len(lines) and indent_of(lines[nxt]) > indent:
                    mapping[key], i = parse(nxt, indent_of(lines[nxt]))
                else:
                    mapping[key], i = None, i + 1

    return parse(0, 0)[0]


# --- GitHub expressions: ${{ a.b == 'x' && 'y' || '' }} -------------------------

TOKEN = re.compile(r"\s*(?:('(?:[^']|'')*')|(\d+(?:\.\d+)?)|(==|!=|&&|\|\||!|\(|\))|([A-Za-z_][\w.-]*))")


def evaluate(expression, context):
    tokens = []
    pos = 0
    expression = expression.strip()
    while pos < len(expression):
        m = TOKEN.match(expression, pos)
        if not m:
            raise ValueError(f"cannot read expression {expression!r} at {pos}")
        pos = m.end()
        string, number, op, name = m.groups()
        if string is not None:
            tokens.append(("lit", string[1:-1].replace("''", "'")))
        elif number is not None:
            tokens.append(("lit", float(number) if "." in number else int(number)))
        elif op is not None:
            tokens.append(("op", op))
        else:
            tokens.append(("lit", {"true": True, "false": False, "null": None}[name]) if name in ("true", "false", "null")
                          else ("name", name))
    tokens.append(("end", None))
    at = [0]

    def peek():
        return tokens[at[0]]

    def take():
        at[0] += 1
        return tokens[at[0] - 1]

    def lookup(name):
        value = context
        for part in name.split("."):
            value = value.get(part) if isinstance(value, dict) else None
        return value

    def primary():
        kind, value = take()
        if kind == "op" and value == "(":
            result = or_()
            take()
            return result
        if kind == "op" and value == "!":
            return not primary()
        if kind == "name":
            if peek() == ("op", "(") and value in ("success", "always"):
                take(), take()
                return True
            return lookup(value)
        return value

    def eq():
        left = primary()
        if peek()[0] == "op" and peek()[1] in ("==", "!="):
            op = take()[1]
            right = primary()
            same = to_text(left).lower() == to_text(right).lower()
            return same if op == "==" else not same
        return left

    def and_():
        left = eq()
        while peek() == ("op", "&&"):
            take()
            right = eq()
            left = right if left else left
        return left

    def or_():
        left = and_()
        while peek() == ("op", "||"):
            take()
            right = and_()
            left = left if left else right
        return left

    return or_()


def to_text(value):
    if value is None:
        return ""
    if value is True or value is False:
        return "true" if value else "false"
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value)


def interpolate(text, context):
    return re.sub(r"\$\{\{(.*?)\}\}", lambda m: to_text(evaluate(m.group(1), context)), text)


class WorkflowRun:
    def __init__(self, sandbox, file, event, inputs, folder):
        self.sandbox, self.event, self.folder = sandbox, event, folder
        spec = parse_yaml((sandbox.work / ".github/workflows" / file).read_text(encoding="utf-8"))
        self.name = spec["name"]
        declared = ((spec.get("on") or {}).get("workflow_dispatch") or {}).get("inputs") or {}
        unknown = set(inputs) - set(declared)
        if unknown:
            raise ValueError(f"{file} has no inputs {unknown}")
        self.inputs = {}
        if event == "workflow_dispatch":
            self.inputs = {k: v.get("default") for k, v in declared.items()}
            self.inputs.update(inputs)
        (self.job,) = spec["jobs"].values()
        self.steps = []          # (name, outcome)
        self.summary = ""
        self.log = ""
        self.conclusion = "success"

    def run(self):
        folder, env0 = self.folder, self.sandbox.env
        folder.mkdir()
        checkout = self.job["steps"][0].get("with") or {}
        repo = folder / "repo"
        depth = [] if str(checkout.get("fetch-depth", "1")) == "0" else ["--depth", "1"]
        run(["git", "clone", "-q", *depth, "--branch", "main", f"file://{self.sandbox.origin}", str(repo)], folder, env0)
        self.sha = run(["git", "rev-parse", "HEAD"], repo, env0).stdout.strip()
        temp = folder / "runner_temp"
        temp.mkdir()
        shims = folder / "bin"
        shims.mkdir()
        for name, body in (("python", f'exec "{sys.executable}" "$@"'), ("pip", "exit 0"),
                           ("gh", 'echo "gh is not available in the test sandbox" >&2; exit 1')):
            (shims / name).write_text(f"#!/bin/bash\n{body}\n")
            (shims / name).chmod(0o755)
        context = {"inputs": self.inputs, "steps": {}, "github": {"token": "test-token", "sha": self.sha}}
        base_env = dict(env0, PATH=f"{shims}{os.pathsep}{env0.get('PATH', '')}", RUNNER_TEMP=str(temp),
                        GITHUB_SHA=self.sha, GITHUB_RUN_ID="1000", GITHUB_SERVER_URL="https://github.com",
                        GITHUB_REPOSITORY="owner/repo", GITHUB_EVENT_NAME=self.event, CI="true")
        for k, v in (self.job.get("env") or {}).items():
            base_env[k] = interpolate(to_text(v), context)
        summary_file = folder / "summary.md"
        summary_file.write_text("")
        failed = False
        for number, step in enumerate(self.job["steps"]):
            name = step.get("name", f"step {number}")
            if "uses" in step:
                self.steps.append((name, "success"))
                continue
            condition = step.get("if")
            if failed or (condition is not None and not evaluate(to_text(condition), context)):
                self.steps.append((name, "skipped"))
                if step.get("id"):
                    context["steps"][step["id"]] = {"outputs": {}, "outcome": "skipped", "conclusion": "skipped"}
                continue
            output_file = folder / f"output-{number}"
            output_file.write_text("")
            env = dict(base_env, GITHUB_OUTPUT=str(output_file), GITHUB_STEP_SUMMARY=str(summary_file))
            for k, v in (step.get("env") or {}).items():
                env[k] = interpolate(to_text(v), context)
            script = folder / f"step-{number}.sh"
            script.write_text(interpolate(step["run"], context))
            result = subprocess.run(["bash", "--noprofile", "--norc", "-eo", "pipefail", str(script)],
                                    cwd=repo, env=env, capture_output=True, text=True)
            self.log += f"\n=== {name} (exit {result.returncode})\n{result.stdout}{result.stderr}"
            outcome = "success" if result.returncode == 0 else "failure"
            outputs = dict(line.split("=", 1) for line in output_file.read_text().splitlines() if "=" in line)
            continue_on_error = step.get("continue-on-error") is True
            if step.get("id"):
                context["steps"][step["id"]] = {
                    "outputs": outputs, "outcome": outcome,
                    "conclusion": "success" if continue_on_error else outcome}
            self.steps.append((name, outcome))
            if outcome == "failure" and not continue_on_error:
                failed = True
                self.conclusion = "failure"
        self.summary = summary_file.read_text()
        self.warnings = re.findall(r"^::warning title=([^:]*)::(.*)$", self.log, re.M)
        return self

    def outcome(self, step_name):
        return dict(self.steps)[step_name]
