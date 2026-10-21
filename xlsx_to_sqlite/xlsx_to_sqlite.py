"""
xlsx_to_sqlite.py

Builds the timeline database (assets/timeline_data.db) from the timeline
workbook plus the list of videos.

  items  - one row per workbook row, plus one row per video
  links  - workbook links (link1..linkN), then a (parent, video) and a
           (video, parent) row for every video, in video id order

The workbook (Sheet1) has these columns, in any order:
    id, era, year, type, hebrew_header, english_header, about_item,
    about_creator, text_historical, text_term, paradigm, paradigm_main,
    paradigm_group, link1, link2, ... linkN

about_creator is stored in the database column about_author. The database
also keeps text_character and text_quote, which are always NULL.

Videos are not in the workbook. They come from a CSV file with the columns
video_id,parent_id; each one must also be in the video manifest and have its
.mp4 and .jpg next to it. A video that is in the manifest but not in the CSV
is only a warning (the site ignores it; Clean up orphans removes it). A
workbook link that points to a video is skipped with a note, because the
build adds video links itself.

With --prune-videos (used by the "Build timeline database" workflow), a
video whose parent was deleted from the workbook is removed from the CSV
instead of stopping the build. A line is removed only when its parent_id is
a whole number, is not an item in the workbook, and was a primary_image in
the --baseline database; any other missing parent (a typo, a blank) is
still a problem. The lines are removed only when the whole build passes,
and the CSV is rewritten with nothing else changed.

Every primary_image and every secondary_image that is not a video needs its
four image files (thumbnails, 200, 1000, 2000) and their entries in the image
manifest.

Everything is validated before anything is written. On any problem, every
problem is printed and nothing is written.

Usage:
    python xlsx_to_sqlite.py WORKBOOK OUTPUT_DB --videos VIDEOS_CSV
        --manifest VIDEO_MANIFEST_JSON --images IMAGE_MANIFEST_JSON
        [--baseline CURRENT_DB] [--allow-shrink]
        [--prune-videos [--prune-summary SUMMARY_MD] [--prune-commit-message MESSAGE_TXT]]

Example (from the repository root):
    python xlsx_to_sqlite/xlsx_to_sqlite.py data/timeline.xlsx new.db \\
        --videos data/videos.csv --manifest assets/videos/manifest.json \\
        --images assets/images/manifest.json --baseline assets/timeline_data.db
"""

import argparse
import csv
import datetime
import json
import os
import re
import sqlite3
import sys
from pathlib import Path

from openpyxl import load_workbook

SHEET_NAME = "Sheet1"

# The site reads exactly these columns; this text matches the live schema.
CREATE_SCHEMA_SQL = """CREATE TABLE items (
    id                    INT,
    era                   TEXT,
    year                  INT,
    type                  TEXT,
    hebrew_header         TEXT,
    english_header        TEXT,
    about_item            TEXT,
    about_author          TEXT,
    text_character        TEXT,
    text_historical       TEXT,
    text_quote            TEXT,
    text_term             TEXT
, paradigm TEXT, paradigm_main TEXT, paradigm_group TEXT);
CREATE TABLE links (
    id        INT,
    target_id INT
);
"""

DB_COLUMNS = [
    "id", "era", "year", "type", "hebrew_header", "english_header",
    "about_item", "about_author", "text_character", "text_historical",
    "text_quote", "text_term", "paradigm", "paradigm_main", "paradigm_group",
]

# Workbook column -> database column, for columns whose names differ.
WORKBOOK_TO_DB = {"about_creator": "about_author"}

REQUIRED_COLUMNS = ["id", "era", "year", "type", "paradigm_group"]
OPTIONAL_COLUMNS = [
    "hebrew_header", "english_header", "about_item", "about_creator",
    "text_historical", "text_term", "paradigm", "paradigm_main",
]
LINK_COLUMN = re.compile(r"^link([1-9][0-9]*)$")

# Kept in the database, never in the workbook.
ALWAYS_NULL_COLUMNS = ["text_character", "text_quote"]

TYPES = ["primary_image", "secondary_image", "historical_event", "tech_event"]

ERAS = [
    "יסודות המחשוב ותפעול ישיר",
    "הממשק הגרפי",
    "מוחשי והיקפי",
    "תפעול במגע ובגוף",
    "טבעי ומרחבי",
    "בינה מלאכותית גנרטיבית והיקפית",
]

PARADIGM_GROUPS = [
    "טבעי וקולי",
    "מוחשי ושולחני",
    "מגע ותחושה",
    "מחשוב שולחני",
    "אינטראקציה עם בינה מלאכותית",
    "נייד ונוכח בכל מקום",
    "מציאות רבודה, מדומה ומרחבית",
    "רובוטים ושליטה מרחוק",
    "היפרטקסט",
    "ביטוי ועיצוב ביקורתי",
    "אינטראקציה לבישה",
]

IMAGE_TYPES = ["primary_image", "secondary_image"]  # secondary_image items that are videos excepted
IMAGE_SIZES = ["thumbnails", "200", "1000", "2000"]

VIDEO_TYPE = "secondary_image"
VIDEO_ENGLISH_HEADER = "Video"
VIDEO_ENGLISH_HEADER_OVERRIDES = {163: "Videoplace (video)"}

# The build fails if it would drop more than this compared with --baseline.
MAX_ITEMS_REMOVED = 5
MAX_LINKS_REMOVED_FRACTION = 0.10


class Problems:
    """Collects every validation problem so they can all be reported at once."""

    def __init__(self):
        self.messages = []

    def add(self, message):
        self.messages.append(message)

    def __bool__(self):
        return bool(self.messages)


def to_int(value):
    """Whole number from an Excel cell (int, 12.0 or "12"), else None."""
    if isinstance(value, bool):
        return None
    if isinstance(value, int):
        return value
    if isinstance(value, float):
        return int(value) if value.is_integer() else None
    if isinstance(value, str) and re.fullmatch(r"\s*-?[0-9]+(\.0+)?\s*", value):
        return int(float(value))
    return None


def describe(value):
    if isinstance(value, (datetime.datetime, datetime.date, datetime.time)):
        return f"{value!r} (Excel turned it into a date; format the column as Number)"
    return repr(value)


def to_text(value):
    """Text cell value, kept exactly as written (no trimming)."""
    if value is None:
        return None
    if isinstance(value, bool):
        return str(value).upper()
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float):
        return str(int(value)) if value.is_integer() else str(value)
    return value


def read_workbook(path, problems):
    """Returns (items, links_by_id) or None if the sheet cannot be read."""
    workbook = load_workbook(path, data_only=True)
    if SHEET_NAME not in workbook.sheetnames:
        problems.add(f"The workbook has no sheet named {SHEET_NAME!r} "
                     f"(sheets: {', '.join(workbook.sheetnames)})")
        return None
    rows = list(workbook[SHEET_NAME].iter_rows(values_only=True))
    if not rows:
        problems.add(f"{SHEET_NAME} is empty")
        return None

    header = [h.strip() if isinstance(h, str) else h for h in rows[0]]
    while header and header[-1] is None:
        header.pop()
    known = set(REQUIRED_COLUMNS) | set(OPTIONAL_COLUMNS)
    link_columns = []
    seen = set()
    for position, name in enumerate(header, start=1):
        if name is None:
            problems.add(f"Header cell {position} is empty")
        elif name in seen:
            problems.add(f"Column {name!r} appears twice")
        elif LINK_COLUMN.match(str(name)):
            link_columns.append(name)
        elif name not in known:
            problems.add(f"Unknown column {name!r}; check its spelling "
                         f"(allowed: {', '.join(REQUIRED_COLUMNS + OPTIONAL_COLUMNS)}, link1..linkN)")
        seen.add(name)
    for name in REQUIRED_COLUMNS:
        if name not in seen:
            problems.add(f"Missing required column {name!r}")
    if problems:
        return None
    for name in OPTIONAL_COLUMNS:
        if name not in seen:
            print(f"note: column {name!r} is not in the workbook; it will be empty for every item")
    link_columns.sort(key=lambda c: int(LINK_COLUMN.match(c).group(1)))
    column = {name: i for i, name in enumerate(header) if name is not None}

    items = []
    links_by_id = {}
    for row_number, row in enumerate(rows[1:], start=2):
        row = list(row) + [None] * (len(header) - len(row))
        if all(v is None for v in row[:len(header)]):
            continue
        cell = lambda name: row[column[name]] if name in column else None

        raw_id = cell("id")
        item_id = to_int(raw_id)
        if item_id is None:
            problems.add(f"Row {row_number}: id is {'blank' if raw_id is None else describe(raw_id)}; "
                         f"it must be a whole number")
            continue
        where = f"Item {item_id} (row {row_number})"

        item = {name: None for name in DB_COLUMNS}
        item["id"] = item_id
        raw_year = cell("year")
        item["year"] = to_int(raw_year)
        if item["year"] is None:
            problems.add(f"{where}: year is {'blank' if raw_year is None else describe(raw_year)}; "
                         f"it must be a whole number")
        for name in ["era", "type", "paradigm_group"] + OPTIONAL_COLUMNS:
            value = cell(name)
            if isinstance(value, (datetime.datetime, datetime.date, datetime.time)):
                problems.add(f"{where}: {name} is {describe(value)}")
            item[WORKBOOK_TO_DB.get(name, name)] = to_text(value)
        items.append((row_number, item))

        targets = []
        for name in link_columns:
            raw = cell(name)
            if raw is None:
                continue
            target = to_int(raw)
            if target is None:
                problems.add(f"{where}: {name} is {describe(raw)}; it must be an item id")
            else:
                targets.append((name, target))
        links_by_id.setdefault(item_id, []).extend(targets)
    return items, links_by_id


def read_videos(path, problems):
    """Returns [(video_id, parent_id)] in video id order."""
    videos = []
    with open(path, newline="", encoding="utf-8-sig") as f:
        reader = csv.DictReader(f)
        if reader.fieldnames is None or [h.strip() for h in reader.fieldnames] != ["video_id", "parent_id"]:
            problems.add(f"{path}: the header must be exactly video_id,parent_id (found {reader.fieldnames})")
            return []
        for line, row in enumerate(reader, start=2):
            row = {k.strip(): v for k, v in row.items()}
            video_id, parent_id = to_int(row["video_id"]), to_int(row["parent_id"])
            if video_id is None or parent_id is None:
                problems.add(f"{path} line {line}: video_id and parent_id must be whole numbers "
                             f"(found {row['video_id']!r}, {row['parent_id']!r})")
                continue
            videos.append((video_id, parent_id))
    seen = set()
    for video_id, _ in videos:
        if video_id in seen:
            problems.add(f"Video {video_id} is listed twice in {path}")
        seen.add(video_id)
    return sorted(videos)


def image_paths(item_id):
    """The four image files (relative to the image manifest's folder) an item needs."""
    return [f"{size}/{item_id}.jpg" for size in IMAGE_SIZES]


def validate(items, links_by_id, videos, manifest_path, images_manifest_path, problems, pruned=None):
    """pruned: the videos --prune-videos removed, or None when not pruning."""
    by_id = {}
    for row_number, item in items:
        if item["id"] in by_id:
            problems.add(f"Item {item['id']}: duplicate id (rows {by_id[item['id']][0]} and {row_number})")
        else:
            by_id[item["id"]] = (row_number, item)

    video_ids = {v for v, _ in videos}
    for row_number, item in items:
        where = f"Item {item['id']} (row {row_number})"
        if item["type"] not in TYPES:
            problems.add(f"{where}: type {item['type']!r} is not one of {', '.join(TYPES)}")
        if item["era"] not in ERAS:
            problems.add(f"{where}: era {item['era']!r} is not one of the six era names")
        group = item["paradigm_group"]
        if item["type"] == "primary_image" and group is None:
            problems.add(f"{where}: primary_image needs a paradigm_group")
        elif group is not None and group not in PARADIGM_GROUPS:
            problems.add(f"{where}: paradigm_group {group!r} is not one of the 11 groups")
        if item["id"] in video_ids:
            problems.add(f"{where}: id {item['id']} is also a video id in the videos list")

    for source_id, targets in links_by_id.items():
        for column, target in targets:
            if target in video_ids:
                continue
            if target not in by_id:
                problems.add(f"Item {source_id}: {column} points to {target}, which is not an item in the workbook")

    manifest_ids = set()
    try:
        manifest = json.loads(Path(manifest_path).read_text(encoding="utf-8"))
        manifest_ids = {entry["id"] for entry in manifest}
    except (OSError, ValueError, KeyError, TypeError) as error:
        problems.add(f"Cannot read the video manifest {manifest_path}: {error}")
    for video_id in sorted(video_ids - manifest_ids):
        problems.add(f"Video {video_id} is in the videos list but not in {manifest_path} "
                     f"(the site would fail to load)")
    for video_id in sorted(manifest_ids - video_ids - {v for v, _ in pruned or []}):
        print(f"warning: video {video_id} is in {manifest_path} but not in the videos list; the site "
              f"ignores it. Run the 'Clean up orphans' workflow to remove its files and manifest entry.")
    video_dir = Path(manifest_path).parent
    for video_id, parent_id in videos:
        for extension in ("mp4", "jpg"):
            if not (video_dir / f"{video_id}.{extension}").is_file():
                problems.add(f"Video {video_id}: file {video_dir / f'{video_id}.{extension}'} is missing")
        parent = by_id.get(parent_id)
        if parent is None and pruned is None:
            problems.add(f"Video {video_id}: parent {parent_id} is not an item in the workbook. If you deleted "
                         f"item {parent_id} on purpose, also delete the line '{video_id},{parent_id}' from "
                         f"data/videos.csv (the build workflow does this by itself)")
        elif parent is None:
            problems.add(f"Video {video_id}: parent {parent_id} is not an item in the workbook, and it was not a "
                         f"picture item before this upload either, so it is not removed automatically. Check "
                         f"the parent id in the line '{video_id},{parent_id}' of data/videos.csv for a typo; to "
                         f"remove the video, delete that line")
        elif parent[1]["type"] != "primary_image":
            problems.add(f"Video {video_id}: parent {parent_id} is a {parent[1]['type']}, not a primary_image")

    image_entries = set()
    try:
        image_manifest = json.loads(Path(images_manifest_path).read_text(encoding="utf-8"))
        image_entries = {entry["path"] for entry in image_manifest}
    except (OSError, ValueError, KeyError, TypeError) as error:
        problems.add(f"Cannot read the image manifest {images_manifest_path}: {error}")
    image_dir = Path(images_manifest_path).parent
    for row_number, item in items:
        if item["type"] not in IMAGE_TYPES or item["id"] in video_ids:
            continue
        missing = []
        for path in image_paths(item["id"]):
            gaps = [what for what, ok in (("file", (image_dir / path).is_file()),
                                          ("manifest entry", path in image_entries)) if not ok]
            if gaps:
                missing.append(f"{path} ({' and '.join(gaps)})")
        if missing:
            problems.add(f"Item {item['id']} (row {row_number}): a {item['type']} needs its images, but these are "
                         f"missing: {', '.join(missing)}. If they were removed when the item was deleted (by the "
                         f"rebuild or by 'Clean up orphans'), bring them back first by reverting that removal "
                         f"commit (see data/README.md, 'Bringing deleted files back')")
    return by_id


def build_rows(items, links_by_id, videos, by_id):
    video_ids = {v for v, _ in videos}
    item_rows = [tuple(item[c] for c in DB_COLUMNS) for _, item in items]
    link_rows = []
    for _, item in items:
        for column, target in links_by_id.get(item["id"], []):
            if target in video_ids:
                print(f"note: item {item['id']} {column} -> {target} is a video; skipped "
                      f"(the build adds video links from the videos list)")
                continue
            link_rows.append((item["id"], target))
    for video_id, parent_id in videos:
        parent = by_id[parent_id][1]
        video = {c: None for c in DB_COLUMNS}
        video.update(id=video_id, era=parent["era"], year=parent["year"], type=VIDEO_TYPE,
                     english_header=VIDEO_ENGLISH_HEADER_OVERRIDES.get(video_id, VIDEO_ENGLISH_HEADER))
        item_rows.append(tuple(video[c] for c in DB_COLUMNS))
    for video_id, parent_id in videos:
        link_rows.append((parent_id, video_id))
        link_rows.append((video_id, parent_id))
    return item_rows, link_rows


def videos_to_prune(videos, items, baseline_path):
    """The (video_id, parent_id) pairs whose parent was a primary_image in the
    baseline database and is not an item in the workbook any more."""
    workbook_ids = {item["id"] for _, item in items}
    with sqlite3.connect(f"file:{baseline_path}?mode=ro", uri=True) as con:
        baseline_parents = {i for (i,) in con.execute("SELECT id FROM items WHERE type = 'primary_image'")}
    return [(v, p) for v, p in videos if p not in workbook_ids and p in baseline_parents]


def remove_csv_lines(path, pruned):
    """Rewrites the videos CSV without the given (video_id, parent_id) lines,
    keeping every other byte (header, line endings, order) as it was."""
    path = Path(path)
    raw = path.read_bytes()
    lines = raw.decode("utf-8-sig").splitlines(keepends=True)
    bom = "\ufeff" if raw.startswith(b"\xef\xbb\xbf") else ""
    gone = set(pruned)
    kept = [lines[0]]
    for line in lines[1:]:
        fields = next(csv.reader([line]), [])
        pair = tuple(to_int(f) for f in fields) if len(fields) == 2 else None
        if pair in gone:
            gone.discard(pair)
        else:
            kept.append(line)
    if gone:
        raise RuntimeError(f"could not find the lines {sorted(gone)} in {path}")
    temp = path.with_name(path.name + ".partial")
    temp.write_bytes((bom + "".join(kept)).encode("utf-8"))
    os.replace(temp, path)


def write_prune_reports(pruned, videos_path, manifest_path, summary_path, message_path):
    video_dir = Path(manifest_path).parent
    lines = ["## Videos removed with their deleted parent", "",
             f"These parents were deleted from the workbook in this upload, so their videos' lines were "
             f"removed from `{videos_path}` (a separate commit). The video files and manifest entries are "
             f"removed by **Leftover files of deleted items** below.", "",
             "| video | parent | line removed | files to remove |", "|---|---|---|---|"]
    for video_id, parent_id in pruned:
        files = ", ".join(f"`{video_dir / f'{video_id}.{ext}'}`" for ext in ("mp4", "jpg"))
        lines.append(f"| {video_id} | {parent_id} | `{video_id},{parent_id}` | {files} and its entry in "
                     f"`{manifest_path}` |")
    lines.append("")
    if summary_path:
        Path(summary_path).write_text("\n".join(lines) + "\n", encoding="utf-8")
    if message_path:
        ids = ", ".join(f"{v} (parent {p})" for v, p in pruned)
        Path(message_path).write_text(
            f"Remove videos of deleted items: {ids}\n\n"
            + "\n".join(f"Removed the line '{v},{p}' from {videos_path}: item {p} was deleted from the workbook."
                        for v, p in pruned) + "\n", encoding="utf-8")


def check_shrink(item_rows, link_rows, baseline_path, allow_shrink, problems):
    with sqlite3.connect(f"file:{baseline_path}?mode=ro", uri=True) as con:
        old_items = con.execute("SELECT COUNT(*) FROM items").fetchone()[0]
        old_links = con.execute("SELECT COUNT(*) FROM links").fetchone()[0]
    removed_items = old_items - len(item_rows)
    removed_links = old_links - len(link_rows)
    print(f"items: {old_items} -> {len(item_rows)}, links: {old_links} -> {len(link_rows)} (compared with {baseline_path})")
    too_many = []
    if removed_items > MAX_ITEMS_REMOVED:
        too_many.append(f"{removed_items} items (limit {MAX_ITEMS_REMOVED})")
    if removed_links > MAX_LINKS_REMOVED_FRACTION * old_links:
        too_many.append(f"{removed_links} links (limit {int(MAX_LINKS_REMOVED_FRACTION * 100)}% = "
                        f"{int(MAX_LINKS_REMOVED_FRACTION * old_links)})")
    if too_many:
        message = "The new database would remove " + " and ".join(too_many)
        if allow_shrink:
            print(f"warning: {message}; allowed by --allow-shrink")
        else:
            problems.add(f"{message}. If that is intended, run again with --allow-shrink "
                         f"(in GitHub: run the workflow manually with allow_shrink ticked)")


def write_database(path, item_rows, link_rows):
    path = Path(path)
    temp = path.with_name(path.name + ".partial")
    temp.unlink(missing_ok=True)
    with sqlite3.connect(temp) as con:
        con.executescript(CREATE_SCHEMA_SQL)
        con.executemany(f"INSERT INTO items VALUES ({', '.join('?' * len(DB_COLUMNS))})", item_rows)
        con.executemany("INSERT INTO links VALUES (?, ?)", link_rows)
    con.close()
    os.replace(temp, path)


def main():
    parser = argparse.ArgumentParser(description="Build the timeline database from the workbook and the videos list.")
    parser.add_argument("workbook", help="the timeline workbook, e.g. data/timeline.xlsx")
    parser.add_argument("output", help="the database file to write")
    parser.add_argument("--videos", required=True, help="CSV with video_id,parent_id, e.g. data/videos.csv")
    parser.add_argument("--manifest", required=True, help="the video manifest, e.g. assets/videos/manifest.json")
    parser.add_argument("--images", required=True, help="the image manifest, e.g. assets/images/manifest.json")
    parser.add_argument("--baseline", help="current database; fail if the new one is much smaller")
    parser.add_argument("--allow-shrink", action="store_true", help="allow a much smaller database")
    parser.add_argument("--prune-videos", action="store_true",
                        help="remove the videos of parents deleted from the workbook from the videos CSV "
                             "(needs --baseline; done only when the whole build passes)")
    parser.add_argument("--prune-summary", help="with --prune-videos: write a Markdown summary here if any were removed")
    parser.add_argument("--prune-commit-message", help="with --prune-videos: write a commit message here if any were removed")
    args = parser.parse_args()
    if args.prune_videos and not args.baseline:
        parser.error("--prune-videos needs --baseline")

    problems = Problems()
    for path in (args.workbook, args.videos, args.manifest, args.images, args.baseline):
        if path and not Path(path).is_file():
            problems.add(f"File not found: {path}")
    if problems:
        return fail(problems)

    workbook = read_workbook(args.workbook, problems)
    videos = read_videos(args.videos, problems)
    if workbook is None:
        return fail(problems)
    items, links_by_id = workbook
    pruned = videos_to_prune(videos, items, args.baseline) if args.prune_videos else []
    videos = [v for v in videos if v not in pruned]
    by_id = validate(items, links_by_id, videos, args.manifest, args.images, problems,
                     pruned if args.prune_videos else None)
    if problems:
        return fail(problems, pruned)

    item_rows, link_rows = build_rows(items, links_by_id, videos, by_id)
    if args.baseline:
        check_shrink(item_rows, link_rows, args.baseline, args.allow_shrink, problems)
    if problems:
        return fail(problems, pruned)

    if pruned:
        remove_csv_lines(args.videos, pruned)
        write_prune_reports(pruned, args.videos, args.manifest, args.prune_summary, args.prune_commit_message)
        for video_id, parent_id in pruned:
            print(f"Removed video {video_id} from {args.videos}: its parent {parent_id} was deleted from the workbook")
    write_database(args.output, item_rows, link_rows)
    print(f"Done -> {args.output}")
    print(f"  items : {len(item_rows)} rows ({len(items)} from the workbook, {len(videos)} videos)")
    print(f"  links : {len(link_rows)} rows")
    return 0


def fail(problems, pruned=()):
    print(f"Build failed: {len(problems.messages)} problem(s). Nothing was written.", file=sys.stderr)
    for message in problems.messages:
        print(f"  - {message}", file=sys.stderr)
    for video_id, parent_id in pruned:
        print(f"note: video {video_id} belongs to item {parent_id}, which was deleted; it will be removed "
              f"automatically once the problems above are fixed", file=sys.stderr)
    return 1


if __name__ == "__main__":
    sys.exit(main())
